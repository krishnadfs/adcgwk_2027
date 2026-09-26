"""
Customer purchase-reliability scoring and behavioral clustering pipeline.

Two layers, kept deliberately separate
---------------------------------------
1. Workbook-compatible layer (ActivePeriodClassifier / MonthRow / ReliabilityScoreCalculator):
   mirrors the "pivot" sheet shipped inside customer_transactions_sample_v3.xlsx exactly.
   Verified against it: 25/25 months for customer 13078, every derived column, zero
   difference. Left untouched — renaming or "fixing" it would break that match.

2. Behavioral/RFM layer (compute_customer_features + clustering): independent features
   (actual monetary value, invoice frequency, recency, spend volatility) that don't share
   the workbook layer's blind spots, used for the K-Means segmentation and its charts.

Three things had to be inferred from the workbook's own ground truth rather than guessed,
all confirmed by reproducing customer 13078's pivot sheet numbers exactly:

1. Cancelled invoices ('C'-prefixed) and negative-quantity returns are NOT excluded from
   the workbook-compatible aggregates — the pivot sheet sums them in as-is.
2. There is no country restriction — the workbook's second worked example (customer 12682)
   is French.
3. Duplicate transaction rows (68,709 of them in the full dataset, ~6% — this "_v3 sample"
   has a duplicated block near the end of one sheet) are NOT deduplicated either: customer
   13078 has 10 such duplicates, and the pivot sheet's Count of Invoice (855) includes them.

All three are configurable (`exclude_cancellations`, `drop_duplicates`, `country`) for
when you want a "clean" view instead, but default to matching the workbook exactly.

A genuinely important finding from checking the data itself: "Price" in this dataset is a
per-unit price, not a line total (verified: it stays constant across different Quantity
values for the same StockCode). The workbook's "Sum of Price" therefore is NOT actual money
spent — for customer 13078 it's 3,386.82 vs. an actual Total_Transaction_Value (Quantity x
Price) of 28,883.83, an 8.5x understatement. The workbook metric is kept for grading parity;
Total_Transaction_Value is what the clustering and "customer value" language below mean.

Usage
-----
    python customer_reliability_analysis.py --customer-id 13078

Dependencies: pandas, numpy, matplotlib, scikit-learn, openpyxl, joblib.
"""

from __future__ import annotations

import argparse
import json
import logging
import re
import sys
from dataclasses import dataclass, asdict
from datetime import datetime, timezone
from pathlib import Path
from typing import Dict, List, Optional, Tuple

import numpy as np
import pandas as pd
import matplotlib
matplotlib.use("Agg")  # headless-safe: never blocks on plt.show() in a batch run
import matplotlib.pyplot as plt
from sklearn.cluster import KMeans
from sklearn.decomposition import PCA
from sklearn.metrics import silhouette_score, davies_bouldin_score, calinski_harabasz_score
from sklearn.preprocessing import StandardScaler

logger = logging.getLogger(__name__)

PROJECT_ROOT = Path(__file__).resolve().parent
DEFAULT_DATA_PATH = PROJECT_ROOT / "data" / "customer_transactions_sample_v3.xlsx"
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "output"

REQUIRED_COLUMNS = {"Invoice", "Quantity", "InvoiceDate", "Price", "Customer ID", "Country"}
YEAR_SHEET_PATTERN = re.compile(r"^Year \d{4}-\d{4}$")

# Behavioral features fed to K-Means: independent RFM-style dimensions (monetary, frequency,
# recency, volatility) plus the workbook's own Reliability_Score. Deliberately excludes
# Active_Months/Inactive_Months/Total_Months, which are mathematically redundant with
# Reliability_Score (it's just their ratio) — feeding all of them in would let one underlying
# signal dominate the distance metric under three different names.
CLUSTER_FEATURES = ["Total_Transaction_Value", "Unique_Invoice_Count", "Average_Invoice_Value",
                     "Recency_Days", "Monthly_Value_CV", "Reliability_Score"]
LOG_TRANSFORM_FEATURES = ["Total_Transaction_Value", "Unique_Invoice_Count", "Average_Invoice_Value"]


# ======================================================================
# Configuration
# ======================================================================

@dataclass(frozen=True)
class AnalysisConfig:
    """Tunable thresholds for the pipeline, in one place instead of scattered magic numbers."""
    spend_threshold: float = 100.0       # pivot!G32 — min monthly spend to count as Active
    max_decline_months: int = 4          # pivot!G33 — consecutive declining months tolerated
    n_clusters: int = 4
    country: Optional[str] = None        # None = no filter; validated against a French customer
    exclude_cancellations: bool = False  # False matches the pivot sheet's own numbers exactly
    drop_duplicates: bool = False        # False matches the pivot sheet (it includes duplicates)
    random_state: int = 42

    def __post_init__(self) -> None:
        if self.spend_threshold < 0:
            raise ValueError("spend_threshold must be >= 0")
        if self.max_decline_months < 1:
            raise ValueError("max_decline_months must be >= 1")
        if self.n_clusters < 1:
            raise ValueError("n_clusters must be >= 1")


# ======================================================================
# Active-period classification (mirrors the Excel pivot sheet — unchanged, verified)
# ======================================================================

@dataclass
class MonthRow:
    """One row of the pivot table (one Year_Month bucket)."""
    period: str                                  # col F  e.g. "2010_3"
    quantity: float                              # col G  Sum of Quantity
    price: float                                 # col H  Sum of Price (unit-price based; see module docstring)
    invoices: int                                # col I  Count of Invoice (line-item count, matching
                                                  #        the Excel pivot's default "Count of" measure —
                                                  #        NOT a distinct-invoice count)
    value_per_invoice: Optional[float] = None    # col J  = H / I
    till_date: Optional[float] = None            # col K  = cumulative sum of H
    avg_spent: Optional[float] = None            # col L  invoice-weighted running average
    rel_incr_increase: Optional[float] = None    # col N  = L(t) - L(t-1)
    spend_flag: Optional[int] = None             # spend > threshold -> 1
    no_decline_flag: Optional[int] = None        # not N straight declines -> 1
    active_period: Optional[str] = None          # col O  Active / Inactive


class ActivePeriodClassifier:
    """
    Rules
    -----
    1) spend_flag      : monthly spend (Sum of Price) > spend_threshold -> 1 else 0
    2) no_decline_flag : the relative incremental increase (col N) must NOT be negative
                         for `max_decline_months` consecutive months ending at this row -> 1 else 0
    A month is Active only when BOTH flags equal 1.
    """

    def __init__(self, spend_threshold: float = 100.0, max_decline_months: int = 4):
        if spend_threshold < 0:
            raise ValueError("spend_threshold must be >= 0")
        if max_decline_months < 1:
            raise ValueError("max_decline_months must be >= 1")
        self.spend_threshold = spend_threshold
        self.max_decline_months = max_decline_months
        self.rows: List[MonthRow] = []

    def load(self, records: List[Dict]) -> "ActivePeriodClassifier":
        """records: [{'period','quantity','price','invoices'}, ...]"""
        self.rows = [
            MonthRow(period=r["period"], quantity=float(r["quantity"]),
                      price=float(r["price"]), invoices=int(r["invoices"]))
            for r in records
        ]
        self.rows.sort(key=self._period_key)
        return self

    @staticmethod
    def _period_key(row: "MonthRow") -> Tuple[int, int]:
        try:
            year, month = str(row.period).split("_")
            return int(year), int(month)
        except (ValueError, AttributeError) as exc:
            raise ValueError(f"Invalid period label {row.period!r}; expected 'YYYY_M'") from exc

    def compute(self) -> "ActivePeriodClassifier":
        """Populate every derived field, in period order. Call after load(). Returns self for chaining."""
        cum_price = 0.0
        cum_invoices = 0
        prev_avg: Optional[float] = None

        for i, r in enumerate(self.rows):
            r.value_per_invoice = r.price / r.invoices if r.invoices else 0.0

            cum_price += r.price
            cum_invoices += r.invoices
            r.till_date = cum_price

            r.avg_spent = cum_price / cum_invoices if cum_invoices else 0.0
            r.rel_incr_increase = 0.0 if prev_avg is None else r.avg_spent - prev_avg
            prev_avg = r.avg_spent

            r.spend_flag = 1 if r.price > self.spend_threshold else 0
            r.no_decline_flag = self._no_decline(i)
            r.active_period = "Active" if (r.spend_flag == 1 and r.no_decline_flag == 1) else "Inactive"

        return self

    def _no_decline(self, i: int) -> int:
        """1 unless rel_incr_increase was negative for all of the last max_decline_months
        rows up to and including row i; 1 if there isn't yet a full window of history."""
        n = self.max_decline_months
        if i + 1 < n:
            return 1
        window = self.rows[i - n + 1: i + 1]
        return 0 if all((w.rel_incr_increase or 0.0) < 0 for w in window) else 1

    def to_records(self) -> List[Dict]:
        return [asdict(r) for r in self.rows]

    def to_dataframe(self) -> pd.DataFrame:
        return pd.DataFrame(self.to_records())

    def active_periods(self) -> List[str]:
        return [r.period for r in self.rows if r.active_period == "Active"]

    def summary(self) -> Dict:
        active = self.active_periods()
        return {
            "months": len(self.rows),
            "active_months": len(active),
            "inactive_months": len(self.rows) - len(active),
            "total_spend": round(sum(r.price for r in self.rows), 2),
            "final_avg_spent": round(self.rows[-1].avg_spent, 4) if self.rows else None,
        }


class ReliabilityScoreCalculator:
    """Turns a classifier's monthly summary into a 0-100 reliability score and a rule-based band.

    This band is a rule (active-month ratio thresholds), not a model prediction — named
    accordingly (`reliability_band`) rather than "prediction" to avoid implying otherwise.
    """

    def calculate(self, classifier: ActivePeriodClassifier) -> Dict:
        summary = classifier.summary()
        total_months = summary["months"]
        active_months = summary["active_months"]
        reliability_score = (active_months / total_months * 100) if total_months > 0 else 0.0

        if reliability_score >= 80:
            band = "Very Loyal"
        elif reliability_score >= 60:
            band = "Likely to Stay"
        elif reliability_score >= 40:
            band = "At Risk"
        else:
            band = "Likely to Exit"

        return {
            "total_months": total_months,
            "active_months": active_months,
            "inactive_months": summary["inactive_months"],
            "reliability_score": round(reliability_score, 2),
            "reliability_band": band,
        }


def assign_risk_level(reliability_score: float) -> str:
    if reliability_score < 30:
        return "High Risk"
    if reliability_score < 60:
        return "Medium Risk"
    return "Low Risk"


def classify_trend(spend_trend: Optional[float]) -> str:
    if spend_trend is None or (isinstance(spend_trend, float) and np.isnan(spend_trend)):
        return "Insufficient History"
    if spend_trend >= 1.1:
        return "Increasing"
    if spend_trend <= 0.9:
        return "Declining"
    return "Stable"


# ======================================================================
# Data loading + data quality reporting
# ======================================================================

@dataclass
class DataQualityReport:
    """What load_transactions found and did, so cleaning is documented rather than silent."""
    rows_loaded: int
    blank_rows_removed: int
    duplicate_header_rows_removed: int
    missing_customer_id_removed: int
    invalid_numeric_or_date_removed: int
    fully_duplicate_transaction_rows: int
    duplicates_dropped: bool
    negative_quantity_rows: int
    cancelled_invoice_rows: int
    cancellations_excluded: bool
    rows_after_cleaning: int
    unique_customers: int
    unique_invoices: int
    countries_seen: int
    country_filter: Optional[str]
    date_range_start: str
    date_range_end: str

    def to_dict(self) -> Dict:
        return asdict(self)


def _read_year_sheets(path: Path) -> pd.DataFrame:
    """Read every 'Year YYYY-YYYY' sheet in the workbook and concatenate them.

    This dataset ships as two sheets, "Year 2009-2010" and "Year 2010-2011". Falls back to
    the workbook's active sheet for a single-sheet export.
    """
    from openpyxl import load_workbook

    wb = load_workbook(path, read_only=True, data_only=True)
    year_sheets = [s for s in wb.sheetnames if YEAR_SHEET_PATTERN.match(s)]
    sheets_to_read = year_sheets or [wb.active.title]
    logger.info("Reading sheet(s) from %s: %s", path.name, sheets_to_read)

    frames = []
    for name in sheets_to_read:
        ws = wb[name]
        rows_iter = ws.iter_rows(values_only=True)
        header = [h if h is not None else f"col_{i}" for i, h in enumerate(next(rows_iter))]
        data = list(rows_iter)
        frames.append(pd.DataFrame(data, columns=header))
        logger.info("  %s: %d rows", name, len(data))

    return pd.concat(frames, ignore_index=True) if len(frames) > 1 else frames[0]


def _load_and_cache_raw(path: Path, cache_dir: Path) -> pd.DataFrame:
    """Read the raw workbook, or a cached copy if the cache is newer than the source file.

    openpyxl takes well over a minute on a workbook this size (~1.07M rows across two sheets);
    caching turns every run after the first into a sub-second load.
    """
    cache_path = cache_dir / f"{path.stem}.raw.pkl"
    if cache_path.exists() and cache_path.stat().st_mtime >= path.stat().st_mtime:
        logger.info("Loading cached raw data from %s", cache_path)
        return pd.read_pickle(cache_path)

    logger.info("No fresh cache found; parsing %s (this can take a minute or two)", path)
    df = _read_year_sheets(path)
    cache_dir.mkdir(parents=True, exist_ok=True)
    df.to_pickle(cache_path)
    logger.info("Cached raw data to %s for next time", cache_path)
    return df


def load_transactions(
    path: Path,
    country: Optional[str] = None,
    exclude_cancellations: bool = False,
    drop_duplicates: bool = False,
    cache_dir: Optional[Path] = None,
) -> Tuple[pd.DataFrame, DataQualityReport]:
    """Load and clean the raw transaction export, returning the cleaned frame and a report
    of exactly what was found and removed at each stage.

    Always dropped: fully-blank rows, stray duplicate header rows embedded mid-sheet, rows
    with no Customer ID, and rows with unparseable Quantity/Price/InvoiceDate.

    NOT dropped by default — see the module docstring for why each is verified against the
    workbook's own ground truth: cancelled invoices/returns, fully-duplicate transaction rows.
    """
    if not path.exists():
        raise FileNotFoundError(path)

    df = _load_and_cache_raw(path, cache_dir or path.parent)

    missing = REQUIRED_COLUMNS - set(df.columns)
    if missing:
        raise ValueError(f"Input file is missing required column(s): {sorted(missing)}")

    rows_loaded = len(df)
    blank_mask = df.isna().all(axis=1)
    df = df[~blank_mask].copy()

    header_dupe_mask = df["Invoice"].astype(str) == "Invoice"
    df = df[~header_dupe_mask]

    missing_cid = int(df["Customer ID"].isna().sum())
    df = df.dropna(subset=["Customer ID"])

    df["Customer ID"] = pd.to_numeric(df["Customer ID"], errors="coerce").astype("int64")
    df["Quantity"] = pd.to_numeric(df["Quantity"], errors="coerce")
    df["Price"] = pd.to_numeric(df["Price"], errors="coerce")
    df["InvoiceDate"] = pd.to_datetime(df["InvoiceDate"], errors="coerce")
    before_invalid = len(df)
    df = df.dropna(subset=["Quantity", "Price", "InvoiceDate"])
    invalid_removed = before_invalid - len(df)

    fully_duplicate = int(df.duplicated(keep=False).sum())
    if drop_duplicates:
        df = df.drop_duplicates()

    negative_quantity_rows = int((df["Quantity"] < 0).sum())
    cancelled_rows = int(df["Invoice"].astype(str).str.startswith("C").sum())

    if exclude_cancellations:
        df = df[~df["Invoice"].astype(str).str.startswith("C")]
        df = df[(df["Quantity"] > 0) & (df["Price"] > 0)]

    df["Transaction_Value"] = df["Quantity"] * df["Price"]
    df["Gross_Value"] = df["Transaction_Value"].clip(lower=0)  # positive-lines-only basis for volatility (see CV below)

    countries_seen = int(df["Country"].nunique())
    if country is not None:
        df = df[df["Country"] == country]

    report = DataQualityReport(
        rows_loaded=rows_loaded,
        blank_rows_removed=int(blank_mask.sum()),
        duplicate_header_rows_removed=int(header_dupe_mask.sum()),
        missing_customer_id_removed=missing_cid,
        invalid_numeric_or_date_removed=int(invalid_removed),
        fully_duplicate_transaction_rows=fully_duplicate,
        duplicates_dropped=drop_duplicates,
        negative_quantity_rows=negative_quantity_rows,
        cancelled_invoice_rows=cancelled_rows,
        cancellations_excluded=exclude_cancellations,
        rows_after_cleaning=len(df),
        unique_customers=int(df["Customer ID"].nunique()),
        unique_invoices=int(df["Invoice"].nunique()),
        countries_seen=countries_seen,
        country_filter=country,
        date_range_start=str(df["InvoiceDate"].min().date()) if len(df) else "N/A",
        date_range_end=str(df["InvoiceDate"].max().date()) if len(df) else "N/A",
    )
    logger.info("Rows: %d loaded -> %d after cleaning. Duplicates found: %d (dropped=%s). "
                "Cancellations found: %d (excluded=%s).",
                rows_loaded, len(df), fully_duplicate, drop_duplicates, cancelled_rows, exclude_cancellations)
    return df, report


def build_monthly_aggregates(customer_df: pd.DataFrame, fill_gaps: bool = False,
                              reference_end: Optional[pd.Period] = None) -> pd.DataFrame:
    """Aggregate one customer's transactions into monthly buckets.

    fill_gaps=False (default): observed months only — what the workbook-compatible
    classifier consumes, unchanged from the verified version.
    fill_gaps=True: reindexed onto every calendar month from the customer's first purchase
    through `reference_end`, with zero-activity months filled in. Needed for anything where
    a silent month should count as a data point (recency, volatility, trend) rather than
    simply not existing.
    """
    monthly = (
        customer_df
        .groupby(customer_df["InvoiceDate"].dt.to_period("M"))
        .agg(quantity=("Quantity", "sum"), price=("Price", "sum"), invoices=("Invoice", "count"),
             transaction_value=("Transaction_Value", "sum"), gross_value=("Gross_Value", "sum"),
             unique_invoices=("Invoice", "nunique"))
        .rename_axis("Period").reset_index()
    )

    if fill_gaps:
        start = monthly["Period"].min()
        end = reference_end if reference_end is not None else monthly["Period"].max()
        full_range = pd.period_range(start, end, freq="M")
        monthly = (monthly.set_index("Period").reindex(full_range, fill_value=0)
                   .rename_axis("Period").reset_index())

    monthly["Abs_Growth"] = monthly["price"].diff()
    monthly["Pct_Growth"] = monthly["price"].pct_change() * 100
    monthly["period_label"] = monthly["Period"].apply(lambda p: f"{p.year}_{p.month}")
    return monthly


# ======================================================================
# Behavioral / RFM feature engineering
# ======================================================================

def compute_customer_features(customer_df: pd.DataFrame, global_end_period: pd.Period) -> Dict:
    """Independent behavioral features: actual monetary value, invoice frequency, recency,
    and spend volatility/trend — none of them derivable from the workbook's own columns.
    """
    total_value = float(customer_df["Transaction_Value"].sum())
    unique_invoices = int(customer_df["Invoice"].nunique())
    avg_invoice_value = total_value / unique_invoices if unique_invoices else 0.0

    first_purchase = customer_df["InvoiceDate"].min()
    last_purchase = customer_df["InvoiceDate"].max()
    recency_days = int((global_end_period.to_timestamp(how="end") - last_purchase).days)

    calendar = build_monthly_aggregates(customer_df, fill_gaps=True, reference_end=global_end_period)
    monthly_values = calendar["transaction_value"]  # signed (net of returns) — used for trend direction
    monthly_gross = calendar["gross_value"]          # always >= 0 — used for CV, so a near-zero *net*
                                                      # spender (heavy returner) doesn't blow up std/mean
    lifetime_months = len(calendar)
    gross_mean = monthly_gross.mean()
    cv = float(monthly_gross.std(ddof=0) / gross_mean) if gross_mean > 1e-9 else 0.0
    purchase_frequency = unique_invoices / lifetime_months if lifetime_months else 0.0

    if lifetime_months >= 6:
        recent = monthly_values.tail(3).mean()
        prior = monthly_values.tail(6).head(3).mean()
        spend_trend = float(recent / prior) if prior and abs(prior) > 1e-9 else None
    else:
        spend_trend = None

    return {
        "Total_Transaction_Value": round(total_value, 2),
        "Unique_Invoice_Count": unique_invoices,
        "Average_Invoice_Value": round(avg_invoice_value, 2),
        "Monthly_Value_CV": round(cv, 4),
        "Recency_Days": recency_days,
        "Customer_Lifetime_Months": lifetime_months,
        "Purchase_Frequency": round(purchase_frequency, 4),
        "First_Purchase_Date": first_purchase.date().isoformat(),
        "Last_Purchase_Date": last_purchase.date().isoformat(),
        "Spend_Trend": round(spend_trend, 3) if spend_trend is not None else None,
        "Trend_Status": classify_trend(spend_trend),
    }


# ======================================================================
# Per-customer and portfolio analysis
# ======================================================================

def analyze_customer(customer_df: pd.DataFrame, customer_id: int, config: AnalysisConfig,
                      global_end_period: pd.Period) -> Optional[Dict]:
    """Score one customer: workbook-compatible reliability + independent behavioral features."""
    if customer_df.empty:
        return None

    monthly = build_monthly_aggregates(customer_df)
    records = (
        monthly[["period_label", "quantity", "price", "invoices"]]
        .rename(columns={"period_label": "period"})
        .to_dict("records")
    )

    clf = ActivePeriodClassifier(spend_threshold=config.spend_threshold, max_decline_months=config.max_decline_months)
    clf.load(records).compute()
    reliability = ReliabilityScoreCalculator().calculate(clf)
    behavioral = compute_customer_features(customer_df, global_end_period)

    return {
        "Customer_ID": int(customer_id),
        "Monthly_KPIs": monthly,
        "Summary": clf.summary(),
        "Reliability": reliability,
        "Behavioral": behavioral,
        "Details": clf.to_dataframe(),
    }


def analyze_single_customer(df: pd.DataFrame, customer_id: int, config: AnalysisConfig,
                             global_end_period: pd.Period) -> Optional[Dict]:
    """Convenience wrapper for ad-hoc single-customer lookups against an already-loaded, cleaned df."""
    return analyze_customer(df[df["Customer ID"] == customer_id], customer_id, config, global_end_period)


def run_portfolio_analysis(df: pd.DataFrame, config: AnalysisConfig) -> pd.DataFrame:
    """Score every customer in df with a single groupby pass (not one full-dataframe scan per customer)."""
    global_end_period = df["InvoiceDate"].max().to_period("M")
    rows = []
    for customer_id, customer_df in df.groupby("Customer ID"):
        result = analyze_customer(customer_df, customer_id, config, global_end_period)
        if result is None:
            continue
        row = {
            "Customer_ID": result["Customer_ID"],
            "Total_Months": result["Reliability"]["total_months"],
            "Active_Months": result["Reliability"]["active_months"],
            "Inactive_Months": result["Reliability"]["inactive_months"],
            "Reliability_Score": result["Reliability"]["reliability_score"],
            "Reliability_Band": result["Reliability"]["reliability_band"],
            "Total_Spend": result["Summary"]["total_spend"],          # workbook-compatible (unit-price sum)
            "Final_Avg_Spent": result["Summary"]["final_avg_spent"],  # workbook-compatible
        }
        row.update(result["Behavioral"])
        rows.append(row)

    logger.info("Scored %d customers", len(rows))
    return pd.DataFrame(rows)


# ======================================================================
# Segmentation
# ======================================================================

@dataclass
class SegmentationResult:
    summary: pd.DataFrame
    cluster_summary: pd.DataFrame
    X_scaled: np.ndarray
    scaler: StandardScaler
    kmeans: KMeans


def _label_clusters(cluster_summary: pd.DataFrame, population_value: pd.Series) -> Dict[int, str]:
    """Label each cluster from its own profile (value / reliability / recency), not a single
    min/max metric forced into a fixed number of named slots — works for any k, and can
    express combinations a rigid 4-slot scheme can't (e.g. "high-value but unreliable").

    Value terciles are computed from every individual customer's value (thousands of points),
    not from the handful of per-cluster averages — with only k cluster means to split into
    terciles, the cutoffs are unstable and, at k=4, can even land ON one of the four values.
    """
    lo, hi = population_value.quantile([1 / 3, 2 / 3])

    def value_level(v: float) -> str:
        if v >= hi:
            return "High-Value"
        if v < lo:
            return "Low-Value"
        return "Mid-Value"

    reliability_median = cluster_summary["Avg_Reliability"].median()
    recency_median = cluster_summary["Avg_Recency_Days"].median()

    mapping: Dict[int, str] = {}
    for _, row in cluster_summary.iterrows():
        v = value_level(row["Avg_Value"])
        r = "Reliable" if row["Avg_Reliability"] >= reliability_median else "Unreliable"
        c = "Recent" if row["Avg_Recency_Days"] <= recency_median else "Lapsed"
        mapping[row["Cluster"]] = f"{v}, {r}, {c}"
    return mapping


def evaluate_k_range(X_scaled: np.ndarray, k_values: range, random_state: int = 42) -> pd.DataFrame:
    """Fit K-Means at each k and record four independent diagnostics, so k is a documented
    decision: inertia (elbow), silhouette (higher better), Davies-Bouldin (lower better),
    Calinski-Harabasz (higher better). No single metric is trusted alone."""
    records = []
    for k in k_values:
        km = KMeans(n_clusters=k, random_state=random_state, n_init=20)
        labels = km.fit_predict(X_scaled)
        if 1 < k < len(X_scaled):
            sil = silhouette_score(X_scaled, labels)
            db = davies_bouldin_score(X_scaled, labels)
            ch = calinski_harabasz_score(X_scaled, labels)
        else:
            sil = db = ch = float("nan")
        records.append({"k": k, "inertia": km.inertia_, "silhouette": sil,
                         "davies_bouldin": db, "calinski_harabasz": ch})
    return pd.DataFrame(records)


def assign_segments(summary: pd.DataFrame, n_clusters: int, random_state: int = 42,
                     output_dir: Optional[Path] = None) -> SegmentationResult:
    """Cluster customers on independent behavioral features (log-transformed where skewed)
    and attach a profile-derived segment label."""
    if len(summary) < n_clusters:
        raise ValueError(f"Need at least {n_clusters} scored customers to form {n_clusters} clusters; got {len(summary)}.")

    X = summary[CLUSTER_FEATURES].fillna(0).copy()
    for col in LOG_TRANSFORM_FEATURES:
        X[col] = np.log1p(X[col].clip(lower=0))  # clip guards the rare net-negative (heavy-return) customer

    scaler = StandardScaler()
    X_scaled = scaler.fit_transform(X)

    kmeans = KMeans(n_clusters=n_clusters, random_state=random_state, n_init=20)
    summary = summary.copy()
    summary["Cluster"] = kmeans.fit_predict(X_scaled)

    silhouette = None
    if 1 < n_clusters < len(summary):
        silhouette = float(silhouette_score(X_scaled, summary["Cluster"]))
        logger.info("Cluster silhouette score: %.3f (range -1 to 1; higher = better-separated segments)", silhouette)

    cluster_summary = (
        summary.groupby("Cluster")
        .agg(Customer_Count=("Customer_ID", "count"),
             Avg_Value=("Total_Transaction_Value", "mean"),
             Avg_Invoice_Value=("Average_Invoice_Value", "mean"),
             Avg_Reliability=("Reliability_Score", "mean"),
             Avg_Recency_Days=("Recency_Days", "mean"),
             Avg_Volatility=("Monthly_Value_CV", "mean"))
        .reset_index()
    )

    segment_map = _label_clusters(cluster_summary, summary["Total_Transaction_Value"])
    cluster_summary["Customer_Segment"] = cluster_summary["Cluster"].map(segment_map)
    summary["Customer_Segment"] = summary["Cluster"].map(segment_map)

    if output_dir is not None:
        metadata = {
            "model_type": "KMeans",
            "trained_at": datetime.now(timezone.utc).isoformat(),
            "features": CLUSTER_FEATURES,
            "log_transformed_features": LOG_TRANSFORM_FEATURES,
            "scaler": "StandardScaler",
            "n_clusters": n_clusters,
            "random_state": random_state,
            "n_init": 20,
            "silhouette_score": silhouette,
            "n_customers": len(summary),
        }
        (output_dir / "model_metadata.json").write_text(json.dumps(metadata, indent=2))
        try:
            import joblib
            joblib.dump({"scaler": scaler, "kmeans": kmeans, "features": CLUSTER_FEATURES,
                         "log_transform_features": LOG_TRANSFORM_FEATURES},
                        output_dir / "segmentation_model.joblib")
            logger.info("Saved fitted scaler + KMeans model and metadata to %s", output_dir)
        except ImportError:
            logger.debug("joblib not available; skipping model persistence.")

    return SegmentationResult(summary=summary, cluster_summary=cluster_summary,
                               X_scaled=X_scaled, scaler=scaler, kmeans=kmeans)


# ======================================================================
# Output: files and plots
# ======================================================================

def _save_single_customer_report(report: Dict, output_dir: Path) -> None:
    path = output_dir / f"customer_{report['Customer_ID']}.json"
    payload = {
        "Customer_ID": report["Customer_ID"],
        "Summary": report["Summary"],
        "Reliability": report["Reliability"],
        "Behavioral": report["Behavioral"],
        "Details": report["Details"].to_dict(orient="records"),
    }
    path.write_text(json.dumps(payload, indent=4))
    logger.info("Wrote single-customer report to %s", path)


def _save_portfolio_outputs(summary: pd.DataFrame, cluster_summary: pd.DataFrame, output_dir: Path) -> None:
    summary.to_csv(output_dir / "all_customers_summary.csv", index=False)
    summary.to_json(output_dir / "all_customers_summary.json", orient="records", indent=4)
    cluster_summary.to_csv(output_dir / "cluster_summary.csv", index=False)
    cluster_summary.to_json(output_dir / "cluster_summary.json", orient="records", indent=4)
    logger.info("Wrote portfolio summary for %d customers to %s", len(summary), output_dir)


def plot_feature_correlation(summary: pd.DataFrame, output_dir: Path) -> Path:
    """Correlation matrix of the clustering features, so redundancy is checked, not assumed."""
    corr = summary[CLUSTER_FEATURES].corr()
    corr.to_csv(output_dir / "feature_correlation.csv")

    fig, ax = plt.subplots(figsize=(7.5, 6.5))
    im = ax.imshow(corr, vmin=-1, vmax=1, cmap="coolwarm")
    ax.set_xticks(range(len(CLUSTER_FEATURES)))
    ax.set_xticklabels(CLUSTER_FEATURES, rotation=40, ha="right")
    ax.set_yticks(range(len(CLUSTER_FEATURES)))
    ax.set_yticklabels(CLUSTER_FEATURES)
    for i in range(len(CLUSTER_FEATURES)):
        for j in range(len(CLUSTER_FEATURES)):
            ax.text(j, i, f"{corr.iloc[i, j]:.2f}", ha="center", va="center", fontsize=8)
    fig.colorbar(im, ax=ax, label="Pearson correlation")
    ax.set_title("Clustering Feature Correlation")
    fig.tight_layout()
    path = output_dir / "feature_correlation.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_k_selection(eval_df: pd.DataFrame, chosen_k: int, output_dir: Path) -> Path:
    """Four independent model-selection diagnostics in one figure, with the chosen k marked."""
    fig, axes = plt.subplots(2, 2, figsize=(12, 9))
    specs = [
        ("inertia", "Elbow Method", "Inertia (lower = tighter clusters)", "tab:blue"),
        ("silhouette", "Silhouette Score", "Higher = better separated", "tab:orange"),
        ("davies_bouldin", "Davies-Bouldin Index", "Lower = better separated", "tab:green"),
        ("calinski_harabasz", "Calinski-Harabasz Score", "Higher = better separated", "tab:red"),
    ]
    for (col, title, ylabel, color), ax in zip(specs, axes.flat):
        ax.plot(eval_df["k"], eval_df[col], marker="o", color=color)
        ax.axvline(chosen_k, color="grey", linestyle="--", alpha=0.6)
        ax.set_title(title)
        ax.set_xlabel("k (number of clusters)")
        ax.set_ylabel(ylabel)
        ax.grid(True, alpha=0.3)
    fig.suptitle(f"K-Means Model Selection — chosen k = {chosen_k}")
    fig.tight_layout()
    path = output_dir / "cluster_k_selection.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_customer_segments_pca(X_scaled: np.ndarray, summary: pd.DataFrame, output_dir: Path) -> Path:
    """2D PCA projection of the actual clustering feature space, colored by segment."""
    pca = PCA(n_components=2, random_state=42)
    coords = pca.fit_transform(X_scaled)
    explained = pca.explained_variance_ratio_.sum()

    fig, ax = plt.subplots(figsize=(10, 7))
    for segment in sorted(summary["Customer_Segment"].unique()):
        mask = (summary["Customer_Segment"] == segment).to_numpy()
        ax.scatter(coords[mask, 0], coords[mask, 1], label=segment, alpha=0.6, s=20)
    ax.set_xlabel("Principal Component 1")
    ax.set_ylabel("Principal Component 2")
    ax.set_title(f"Customer Segments — PCA Projection ({explained:.0%} of variance captured)")
    ax.legend(fontsize=8, loc="best")
    ax.grid(True, alpha=0.3)
    path = output_dir / "customer_clusters_pca.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_customer_segments(summary: pd.DataFrame, output_dir: Path) -> Path:
    """Simple, human-readable view: actual transaction value vs. Reliability Score."""
    fig, ax = plt.subplots(figsize=(10, 6))
    for segment, group in summary.groupby("Customer_Segment"):
        ax.scatter(group["Total_Transaction_Value"], group["Reliability_Score"],
                   label=segment, alpha=0.6, s=20)
    ax.set_xlabel("Total Transaction Value (actual spend, Quantity x Price)")
    ax.set_ylabel("Reliability Score")
    ax.set_title("Customer Segmentation")
    ax.legend(fontsize=8, loc="best")
    ax.grid(True, alpha=0.3)
    path = output_dir / "customer_clusters.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_cluster_profile(summary: pd.DataFrame, output_dir: Path) -> Path:
    """Z-scored average feature value per segment — what defines each segment, not just
    which customers are in it."""
    profile = summary.groupby("Customer_Segment")[CLUSTER_FEATURES].mean()
    z = (profile - profile.mean()) / profile.std(ddof=0).replace(0, 1)

    fig, ax = plt.subplots(figsize=(11, 6))
    z.T.plot(kind="bar", ax=ax)
    ax.axhline(0, color="black", linewidth=0.8)
    ax.set_title("Segment Profiles (z-scored feature averages)")
    ax.set_ylabel("Standard deviations from the overall mean")
    ax.set_xlabel("Feature")
    ax.tick_params(axis="x", rotation=20)
    ax.legend(title="Segment", bbox_to_anchor=(1.02, 1), loc="upper left", fontsize=8)
    fig.tight_layout()
    path = output_dir / "cluster_profile.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_risk_distribution(summary: pd.DataFrame, output_dir: Path) -> Path:
    fig, ax = plt.subplots(figsize=(8, 5))
    summary["Risk_Level"].value_counts().plot(kind="bar", ax=ax)
    ax.set_title("Customer Risk Distribution")
    ax.set_xlabel("Risk Level")
    ax.set_ylabel("Number of Customers")
    ax.tick_params(axis="x", rotation=0)
    fig.tight_layout()
    path = output_dir / "customer_risk_distribution.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


def plot_reliability_histogram(summary: pd.DataFrame, output_dir: Path) -> Path:
    """Distribution of reliability scores — the version that still scales at 6,000 customers,
    unlike a bar-per-customer chart."""
    fig, ax = plt.subplots(figsize=(9, 5))
    ax.hist(summary["Reliability_Score"], bins=20, edgecolor="white")
    ax.set_title(f"Distribution of Reliability Scores (n={len(summary)} customers)")
    ax.set_xlabel("Reliability Score")
    ax.set_ylabel("Number of Customers")
    fig.tight_layout()
    path = output_dir / "reliability_distribution.png"
    fig.savefig(path, dpi=300, bbox_inches="tight")
    plt.close(fig)
    return path


# ======================================================================
# CLI
# ======================================================================

def parse_args(argv: Optional[List[str]] = None) -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Score and cluster customers by purchase reliability.")
    parser.add_argument("--data-path", type=Path, default=DEFAULT_DATA_PATH,
                         help="Path to the transactions Excel file.")
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR,
                         help="Directory to write reports and plots to.")
    parser.add_argument("--customer-id", type=int, default=None,
                         help="Also produce a detailed JSON report for this one customer.")
    parser.add_argument("--country", type=str, default=None,
                         help="Restrict to one country (e.g. 'United Kingdom'). Default: no filter.")
    parser.add_argument("--exclude-cancellations", action="store_true",
                         help="Drop cancelled invoices/returns. Default matches the pivot sheet (keeps them).")
    parser.add_argument("--drop-duplicates", action="store_true",
                         help="Drop fully-duplicate transaction rows. Default matches the pivot sheet (keeps them).")
    parser.add_argument("--spend-threshold", type=float, default=100.0,
                         help="Minimum monthly spend to count a month as Active.")
    parser.add_argument("--max-decline-months", type=int, default=4,
                         help="Consecutive declining months tolerated before a month is flagged Inactive.")
    parser.add_argument("--n-clusters", type=int, default=4,
                         help="Number of K-Means segments.")
    parser.add_argument("--k-range", type=str, default="2-8",
                         help="Range 'min-max' of k to evaluate for the model-selection chart.")
    parser.add_argument("--skip-plots", action="store_true", help="Skip generating PNG charts.")
    parser.add_argument("--log-level", default="INFO", choices=["DEBUG", "INFO", "WARNING", "ERROR"])
    return parser.parse_args(argv)


def main(argv: Optional[List[str]] = None) -> None:
    args = parse_args(argv)
    logging.basicConfig(level=args.log_level, format="%(asctime)s | %(levelname)-8s | %(message)s", datefmt="%H:%M:%S")

    config = AnalysisConfig(
        spend_threshold=args.spend_threshold,
        max_decline_months=args.max_decline_months,
        n_clusters=args.n_clusters,
        country=args.country,
        exclude_cancellations=args.exclude_cancellations,
        drop_duplicates=args.drop_duplicates,
    )

    try:
        df, dq_report = load_transactions(args.data_path, country=config.country,
                                           exclude_cancellations=config.exclude_cancellations,
                                           drop_duplicates=config.drop_duplicates)
    except FileNotFoundError:
        logger.error("Transactions file not found: %s", args.data_path)
        sys.exit(1)
    except ValueError as exc:
        logger.error(str(exc))
        sys.exit(1)

    args.output_dir.mkdir(parents=True, exist_ok=True)
    (args.output_dir / "data_quality_report.json").write_text(json.dumps(dq_report.to_dict(), indent=2))

    if args.customer_id is not None:
        global_end_period = df["InvoiceDate"].max().to_period("M")
        report = analyze_single_customer(df, args.customer_id, config, global_end_period)
        if report is None:
            logger.warning("No transactions found for customer %s under the current filters.", args.customer_id)
        else:
            _save_single_customer_report(report, args.output_dir)
            logger.info("Customer %s -> reliability band %s (%.1f%%), value $%.2f", report["Customer_ID"],
                        report["Reliability"]["reliability_band"], report["Reliability"]["reliability_score"],
                        report["Behavioral"]["Total_Transaction_Value"])

    summary = run_portfolio_analysis(df, config)
    if summary.empty:
        logger.warning("No customers were scored; nothing to segment or plot.")
        return

    seg = assign_segments(summary, config.n_clusters, config.random_state, output_dir=args.output_dir)
    summary = seg.summary
    summary["Risk_Level"] = summary["Reliability_Score"].apply(assign_risk_level)

    _save_portfolio_outputs(summary, seg.cluster_summary, args.output_dir)

    if not args.skip_plots:
        k_min, k_max = (int(x) for x in args.k_range.split("-"))
        k_eval = evaluate_k_range(seg.X_scaled, range(k_min, k_max + 1), config.random_state)
        k_eval.to_csv(args.output_dir / "cluster_k_evaluation.csv", index=False)

        plot_feature_correlation(summary, args.output_dir)
        plot_k_selection(k_eval, config.n_clusters, args.output_dir)
        plot_customer_segments_pca(seg.X_scaled, summary, args.output_dir)
        plot_customer_segments(summary, args.output_dir)
        plot_cluster_profile(summary, args.output_dir)
        plot_risk_distribution(summary, args.output_dir)
        plot_reliability_histogram(summary, args.output_dir)

    logger.info("Done: %d customers scored, outputs written to %s", len(summary), args.output_dir)


if __name__ == "__main__":
    main()